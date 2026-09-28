#!/usr/bin/env python3
"""Build Google Drive–ready .docx files from GenStudio Brand, Product, and Persona fields."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Inches, Pt, RGBColor

OUT = Path(__file__).resolve().parent


def style_doc(doc: Document) -> None:
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(11)
    styles["Normal"].font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
    pf = styles["Normal"].paragraph_format
    pf.space_after = Pt(8)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    for i in range(1, 4):
        h = styles[f"Heading {i}"]
        h.font.color.rgb = RGBColor(0x15, 0x15, 0x15)
        h.font.name = "Calibri"


def add_intro(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text)
    p.runs[0].italic = True
    p.runs[0].font.size = Pt(10)
    p.runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)


def bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def para(doc: Document, text: str) -> None:
    doc.add_paragraph(text)


def field_label(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(12)


def build_brand() -> None:
    doc = Document()
    style_doc(doc)
    heading(doc, "Red Hat GenStudio Brand guidelines")
    add_intro(
        doc,
        "Paste these sections into Adobe GenStudio for Performance Marketing: "
        "Brand → Tone of voice, Brand values, Editorial guidelines, Editorial restrictions, "
        "and Channel guidelines → Email. Adobe caps Tone of voice and Brand values at 3–6 lines each. "
        "Do not upload this file’s intro sentences into those fields.",
    )

    heading(doc, "Tone of voice", 2)
    add_intro(doc, "Brand → Tone of voice. Adobe cap: 3–6 guidelines.")
    bullets(
        doc,
        [
            "Focus on stability, enterprise-grade reliability, and solving tangible hybrid cloud engineering challenges.",
            "Emphasize community-driven innovation, flexibility, and freedom from vendor lock-in.",
            "Be clear, direct, and practical. Prefer specific outcomes over slogans.",
            "Acknowledge the reader’s operational problem, then move to useful guidance. Do not open with sympathy or “You know…”",
            "Respect technical audiences. Speak clearly without fluff or overly dramatic marketing hype.",
            "Frame Red Hat as an enabler—giving developers and architects the tools to build anywhere.",
        ],
    )

    heading(doc, "Brand values", 2)
    add_intro(doc, "Brand → Brand values. Adobe cap: 3–6 guidelines.")
    bullets(
        doc,
        [
            "Open: Prefer transparency and collaboration over secrecy and control. Use inclusive language. Do not lecture.",
            "Authentic: Prefer truth over hype. Acknowledge complexity. Do not overpromise.",
            "Helpful: Prefer customer success over company promotion. Give an actionable next step the reader can take.",
            "Brave: Take a clear, principled position. Challenge industry assumptions with respect. Do not attack competitors.",
        ],
    )

    heading(doc, "Editorial guidelines", 2)
    add_intro(doc, "Brand → Editorial guidelines. Adobe recommends 5–10 positive, actionable lines.")
    bullets(
        doc,
        [
            "Refer to products by full name on first use in body copy. Approved abbreviations (e.g., RHEL) are allowed in subject lines/headlines.",
            "Use parentheses only for approved initialisms such as RHEL; for other products, shorten to the approved name without inventing an initialism.",
            "Leave common terms such as AI, HTML, and PIN abbreviated; they do not need expansion on first use.",
            "Refer to Red Hat as “it” (third person) or “we/our” (first-person brand voice). Do not use “they.”",
        ],
    )

    heading(doc, "Editorial restrictions", 2)
    add_intro(doc, "Brand → Editorial restrictions. Adobe recommends 5–10 “Avoid…” lines.")
    bullets(
        doc,
        [
            "Avoid absolute and superlative claims such as leading, best, fastest, only, or 100% unless a citation appears in the copy.",
            "Avoid calling a product or environment “secure” or “more secure.” Name the capability or outcome instead (for example, compliance baselines, live kernel patching).",
            "Avoid AI-typical wording: delve, harness, pivotal, landscape, testament, tapestry, unwavering, and openings like “In today’s fast-paced…”",
            "Avoid vague words used without a direct object: robust, powerful, strong, streamline, leverage, utilize, seamless, frictionless.",
            "Avoid “the” before product names (write “Red Hat Enterprise Linux,” not “the Red Hat Enterprise Linux”).",
            "Avoid “the cloud” as an unqualified noun. Specify whose environment (for example, a public cloud, a hybrid cloud, or an organization’s cloud). Prefer “cloud” as a modifier: cloud workloads, cloud providers.",
            "Avoid “the edge” unless the first use is qualified (for example, the edge of the network). Use “an edge device” for hardware. Use “edge computing” for the concept.",
        ],
    )

    heading(doc, "Channel guidelines — Email", 2)
    add_intro(
        doc,
        "Brand → Channel guidelines → Email. 2–5 lines per field. One job per line. "
        "Do not use banned words in examples.",
    )

    heading(doc, "General", 3)
    bullets(
        doc,
        [
            "Maintain a clear, direct, practical, solution-oriented tone. Use short paragraphs.",
            "Write in active voice. Capitalize product names, acronyms, and proper initialisms.",
            "Contractions are allowed. Avoid flexible and scalable; name the specific behavior instead.",
            "Avoid free, win, and unlock anywhere in the email. Use no-cost for no-charge offers.",
        ],
    )

    heading(doc, "Subject line", 3)
    bullets(
        doc,
        [
            "30–50 characters. Put the offer or outcome in the first 30 characters.",
            "Start with an action verb such as Get, Start, Explore, or Try. Do not use Unlock, Free, or Win.",
            "Align with the main offer in the body. Do not use a trailing period.",
        ],
    )

    heading(doc, "Preheader", 3)
    bullets(
        doc,
        [
            "40–90 characters. Put the extra context in the first 40 characters.",
            "Add the outcome or offer the subject line did not state. Do not repeat the subject verbatim.",
            "Do not use a trailing period.",
        ],
    )

    heading(doc, "Headline", 3)
    bullets(
        doc,
        [
            "Maximum 8 words. Put the primary outcome in the first 3 words.",
            "Lead with the reader outcome; use a product capability only to prove it.",
            "Do not use a trailing period.",
        ],
    )

    heading(doc, "Sub-headline", 3)
    bullets(
        doc,
        [
            "Maximum 45 characters.",
            "Elaborate on the headline promise with a specific outcome or next step.",
            "Do not use a trailing period. Question marks are allowed.",
        ],
    )

    heading(doc, "Body", 3)
    bullets(
        doc,
        [
            "Maximum 3 sentences per section. One idea per sentence. Prioritize outcomes over product description.",
            "Acknowledge a specific operational problem, then give useful guidance. Do not start the first sentence with You.",
            "Keep sentences straightforward. Vary sentence length. Avoid long clauses that delay the point.",
        ],
    )

    heading(doc, "Call-to-action", 3)
    bullets(
        doc,
        [
            "2–4 words. Use action verb + object (Start your trial, Explore RHEL).",
            "Name a specific action. Do not use Learn more, Click here, or Unlock.",
            "Do not use end punctuation.",
        ],
    )

    heading(doc, "How to use this document", 2)
    bullets(
        doc,
        [
            "Upload to Google Drive, then Open with Google Docs if you want to edit online.",
            "Copy each bullet list into the matching GenStudio Brand field. Do not paste this whole document into one field.",
            "Do not add trademark symbols (® / ™) to Brand or Channel lines.",
        ],
    )

    path = OUT / "GenStudio-Brand-guidelines.docx"
    doc.save(path)
    print(f"Wrote {path.name}")


def add_product(doc: Document, name: str, description: str, values: list[str], messaging: list[str]) -> None:
    heading(doc, name, 2)
    field_label(doc, "Description")
    para(doc, description)
    field_label(doc, "Value proposition")
    bullets(doc, values)
    field_label(doc, "Messaging preferences")
    bullets(doc, messaging)


def build_products() -> None:
    doc = Document()
    style_doc(doc)
    heading(doc, "Red Hat GenStudio Products")
    add_intro(
        doc,
        "Paste each product’s Description, Value proposition, and Messaging preferences into "
        "GenStudio → Products. One product per GenStudio Product record. "
        "These fields are injected when the product is selected in Parameters—do not also paste them into email prompts.",
    )

    add_product(
        doc,
        "Red Hat Enterprise Linux",
        "Red Hat® Enterprise Linux® is the consistent operating system for building, deploying, and managing workloads across hybrid cloud. It is for organizations that need a hardened Linux foundation from datacenter to public cloud to the edge of the network.",
        [
            "Provides a hardened foundation as the starting point for protection, with compliance baselines, live kernel patching, and proactive risk reduction across physical, virtual, public cloud, and edge computing locations.",
            "Closes the gap between vulnerability detection and deployment with Lightspeed-assisted triage and remediation, so teams can collapse mean time to remediation (MTTR) from weeks to hours.",
            "Simplifies day-to-day operations with intelligent guidance from Red Hat Lightspeed and trusted expertise, so teams can manage Linux as environments and skill gaps grow.",
            "Gives organizations one consistent operating foundation to run and migrate traditional, containerized, and AI workloads without rebuilding for each platform.",
            "Provides a trusted subscription: expert support, a predictable 10-year major release lifecycle, and 30 years of open source leadership.",
        ],
        [
            "Lead with why Red Hat Enterprise Linux—not why Linux. Assume the audience has already chosen Linux.",
            "Position Red Hat Enterprise Linux as the foundation against AI-accelerated threats, where human-speed patching cycles increase exposure. Name capabilities (compliance baselines, live kernel patching, Lightspeed-assisted detection) and outcomes (lower risk, less unplanned downtime).",
            "When the brief is about vulnerability remediation, frame see it, stage it, ship it: Red Hat Enterprise Linux is the hardened OS, Lightspeed identifies vulnerabilities, Satellite gates patches, and Ansible Automation Platform remediates at scale. Do not lead with Satellite or Ansible, or force that campaign phrase, when Red Hat Enterprise Linux is the only selected product.",
            "Frame the offer as a subscription to confidence, not as support insurance or a one-time download.",
            "Anchor copy on four pillars: trust, protection, simplification, and innovation. Lead with outcomes; use capabilities only to prove those outcomes.",
            "Tell a hybrid cloud story: one operating experience across bare metal, virtual, public cloud, and the edge of the network. Do not frame Red Hat Enterprise Linux as public-cloud-only.",
            "Use Red Hat Lightspeed when the topic is simpler operations, the Linux skills gap, or proactive risk reduction and compliance work. Do not make Lightspeed the entire product story.",
            "Address one audience at a time: Champions get risk mitigation, compliance, operational stability, and a predictable lifecycle. Technical Practitioners / Architects get workload portability, automated patch staging, less manual friction, and closing the gap between exploit detection and deployment.",
            "Use the full product name, Red Hat Enterprise Linux, on first mention; after that, RHEL is acceptable. Never use Enterprise Linux, EL, or Linux as the product name.",
            "Cite proof points in the copy—not in footnotes: 1,400+ certified cloud partners, a 10-year major release lifecycle, 30 years of open source leadership, and collapsing MTTR from weeks to hours. Avoid hype, unproven superlatives, and competitor attacks.",
        ],
    )

    add_product(
        doc,
        "Red Hat product trial",
        "Red Hat product trials provide no-cost, time-limited access to evaluate Red Hat solutions with the same benefits as a paid subscription. They are for organizations and practitioners who need to test a product in their own environment before committing to a purchase—not for production use.",
        [
            "Gives teams no-cost, time-limited access to the same software versions, patches, updates, and Red Hat Customer Portal included with a paid subscription—so evaluation reflects real subscription value, not a limited demo.",
            "Lets organizations test a Red Hat product in their own environment—by download or in a public cloud—and judge real-world fit before committing budget.",
            "Puts a clock on evaluation (most trials last 60 days) so teams move from discovery to a decision instead of an open-ended experiment.",
            "Shortens time to a useful evaluation with get-started guidance, documentation, and a My Trials dashboard that shows how to access the product after activation.",
            "Creates a clear next step from evaluation to a paid subscription when the trial is a fit, including converting trial work to production instances.",
            "Supports hands-on experience and certification prep as well as purchase evaluation—without treating the trial as a production environment.",
        ],
        [
            "Lead with no-cost, time-limited, try-before-you-buy evaluation—never “free,” “win,” or “unlock.”",
            "Emphasize full subscription value: all software versions, patches, updates, and Red Hat Customer Portal—not just the latest code.",
            "Frame as evaluation only: product trials are not for production use, and using them in production violates trial terms.",
            "Most self-serve trials last 60 days and are self-supported; do not promise 24x7 production support unless the specific trial includes it.",
            "Distinguish from the Red Hat Developer program (renewable no-cost membership for development-use) and from pay-as-you-go (not a trial; can run in production).",
            "Address one audience at a time: Champions need proof to justify a purchase; Technical Practitioners / Architects need hands-on testing in their environment; Developers need to validate locally without production or sales language.",
            "Pair the trial with a specific product in multipod emails—the trial is the offer, not a substitute for the product story.",
            "Point to documentation, get-started guides, requirements, and My Trials, then a path to a paid subscription after evaluation.",
            "Use CTAs such as Start your trial, Explore the trial, or Start evaluating; purchase language is the next step after a successful evaluation, not the trial offer itself.",
            "Keep the tone clear, direct, and practical; avoid hype, unproven superlatives, and competitor-trial comparisons.",
            "Use Red Hat product trial on first mention, then product trial or trial; never “free trial.”",
        ],
    )

    add_product(
        doc,
        "Red Hat Developer program",
        "The Red Hat Developer program is a no-cost membership that helps developers build, test, and learn with Red Hat technologies. It is for developers, engineers, architects, and technical team leads who want a trusted place to start building and keep learning.",
        [
            "No-cost access to a renewable one-year subscription to build and test with Red Hat products without buying a production license.",
            "Removes the setup barrier: start in minutes with the Developer Sandbox, trials, and downloads instead of standing up enterprise environments.",
            "Stay current and advance your career with curated learning, expert content, events, and early access to betas—including AI and emerging technologies.",
            "Connects members to a trusted community of builders and Red Hat knowledge in Red Hat Customer Portal.",
        ],
        [
            "Speak peer-to-peer to builders, not enterprise buyers. Tone: practical, inviting, useful.",
            "Lead with joining at no cost and getting started quickly. This is a membership to join, not a platform to purchase.",
            "Do not position as a production subscription, paid support plan, or partner program. Development-use access is the offer.",
            "Treat Red Hat products (such as Red Hat Enterprise Linux or Red Hat OpenShift Platform Plus) as things members can try and learn—not as the product being promoted.",
            "Invite a broad technical audience while keeping developers as the primary reader.",
            "Use action-oriented CTAs such as Join Red Hat Developer, Register for an account, or Start building. Avoid “Buy now,” “Talk to sales,” or other purchase language.",
        ],
    )

    add_product(
        doc,
        "Red Hat OpenShift Platform Plus",
        "Red Hat® OpenShift® Platform Plus is a unified platform for enterprises to build, modernize, and deploy applications at scale. It is for organizations that need one consistent hybrid cloud foundation to bring applications to market across datacenter, public cloud, and the edge of the network.",
        [
            "Unifies application development, operations, governance, and data management on one hybrid cloud platform—so teams no longer assemble separate tools as they scale beyond a few Kubernetes clusters.",
            "Consistent way to build, modernize, and deploy across on-premises, public cloud, and the edge of the network—without trading away quality or governance.",
            "Helps development teams deliver faster, with languages and frameworks they choose, while access controls and policy stay in the workflow.",
            "Enables IT operations and security teams to govern many clusters from one place, reducing operational risk as the environment grows.",
            "Improves developer productivity and time to value; a Forrester Consulting Total Economic Impact™ study commissioned by Red Hat found a composite organization realized 203% three-year ROI, 12-month payback, and $3.1 million in improved developer productivity.",
        ],
        [
            "Speak as the complete, self-managed OpenShift edition for organizations scaling beyond a few clusters—not as a starter Kubernetes offering or a managed cloud service.",
            "Lead with outcomes: faster application delivery with policy and access controls in the workflow, and one consistent experience from datacenter to public cloud to the edge of the network. Use capabilities only to prove those outcomes.",
            "Keep the tone clear, direct, and practical. Avoid hype, unproven superlatives, and competitor attacks.",
            "Tell a hybrid cloud story. Do not frame as public-cloud-only.",
            "Treat access controls and policy as part of the application lifecycle and as something that helps developers move faster—not a bolt-on that slows them down.",
            "Address one audience at a time: business leaders (quality and governance without trade-offs), developers (speed and choice), and IT operations and security teams (governance and lower risk).",
            "Use the full product name, Red Hat OpenShift Platform Plus, on first mention. After that, OpenShift Platform Plus is acceptable; never use OpenShift alone or informal shorthand such as OPP.",
            "When proof points are used, attribute them. Do not present Forrester TEI figures as guaranteed results for every customer.",
        ],
    )

    add_product(
        doc,
        "Red Hat Ansible Automation Platform",
        "Red Hat® Ansible® Automation Platform is an enterprise IT automation solution for building, deploying, and managing end-to-end automation at scale. It is for organizations that need one consistent, governed foundation to automate across datacenter, hybrid cloud, network, and the edge of the network.",
        [
            "Gives organizations one platform to create, run, govern, and measure automation at scale, so teams move from fragmented scripts and homegrown tools to a consistent enterprise practice.",
            "Lets teams automate across datacenter, hybrid cloud, network, and the edge of the network under a single set of processes and policies—without rebuilding for each domain.",
            "Helps more people contribute to automation with certified content, development tools, and a self-service portal, while role-based access keeps work governed as adoption grows.",
            "Responds to changing IT conditions with Event-Driven Ansible, turning monitoring alerts into consistent, repeatable actions instead of manual ticket work.",
            "Speeds playbook creation and day-to-day platform operations with Red Hat Ansible Lightspeed, so teams close the automation skills gap without making AI the entire product story.",
            "Improves return on investment, resilience, and compliance; an IDC Snapshot sponsored by Red Hat found a composite organization realized 668% three-year ROI, 8-month payback, and 61% less unplanned downtime.",
        ],
        [
            "Speak as the enterprise automation platform—not community Ansible, Ansible Core, or AWX—and lead with why Ansible Automation Platform, not why automation.",
            "Lead with outcomes: less unplanned work, faster consistent operations, and visible ROI; use capabilities only to prove those outcomes.",
            "Tell a hybrid story—automate across datacenter, public cloud, network, and the edge of the network—and do not frame as public-cloud-only.",
            "Use Red Hat Ansible Lightspeed when the topic is the automation skills gap, faster playbook creation, or simpler platform administration; do not make Lightspeed the entire product story.",
            "Address one audience at a time: Champions need ROI and risk reduction; Technical Practitioners / Architects need governed execution at scale and Event-Driven responses; Developers need trusted content and faster playbooks.",
            "Use Red Hat Ansible Automation Platform on first mention, then Ansible Automation Platform; never Ansible alone, AAP, or Ansible Tower as the product name.",
            "Keep the tone clear, direct, and practical; avoid hype, unproven superlatives, and competitor attacks.",
            "When proof points are used, attribute the IDC Snapshot sponsored by Red Hat (March 2024) and do not present the figures as guaranteed results for every customer.",
            "In multipod emails, pair the platform with the product being automated (such as Red Hat Enterprise Linux or Red Hat OpenShift Platform Plus); Ansible Automation Platform is the automation layer, not a substitute for that product story.",
        ],
    )

    heading(doc, "How to use this document", 2)
    bullets(
        doc,
        [
            "Upload to Google Drive, then Open with Google Docs if you want to edit online.",
            "Create or update one GenStudio Product per heading. Paste Description, Value proposition, and Messaging preferences into the three product fields.",
            "Do not paste these product fields into the email prompt when the product is selected in Parameters.",
        ],
    )

    path = OUT / "GenStudio-Products.docx"
    doc.save(path)
    print(f"Wrote {path.name}")


def add_persona(
    doc: Document,
    name: str,
    roles: str,
    description: str,
    messaging: list[str],
    alias: str | None = None,
) -> None:
    heading(doc, name, 2)
    field_label(doc, "Roles")
    para(doc, roles)
    field_label(doc, "Description")
    para(doc, description)
    field_label(doc, "Messaging preferences")
    bullets(doc, messaging)
    if alias:
        field_label(doc, "Alias")
        para(doc, alias)


def build_personas() -> None:
    doc = Document()
    style_doc(doc)
    heading(doc, "Red Hat GenStudio Personas")
    add_intro(
        doc,
        "WIP — does not represent all personas in GenStudio. "
        "Paste each persona’s Description and Messaging preferences into GenStudio → Personas. "
        "Source of the working copy: GenStudio Personas_WIP.",
    )

    add_persona(
        doc,
        "Champion",
        "System Administrator, Data Science Lead, Automation Architect, AppDev ITDM, Lead Business Analyst",
        "Driven by a desire to solve critical business problems and deliver clear ROI, the Champion is an internal advocate who builds organizational urgency and aligns stakeholders to drive opportunities forward. Grounded in positive past experiences, this group relies on deep vendor trust and needs ready access to responsive sales and technical support to navigate the buying process.",
        [
            "Strategic Growth & Innovation (Forward-Looking Vision) — Seeks forward-looking insights to justify strategic initiatives, showing leadership how modernizing their tech stack drives long-term competitive advantage and business growth.",
            "Operational Impact & Team Empowerment (Solution-Oriented Value) — Needs clear messaging that directly connects tool capabilities to removing operational delays, reducing burnout, and improving overall team efficiency.",
            "Proof & Internal Enablement (Validation & Business Justification) — Requires concrete functional proof, reliable performance data, and clear financial justification to create internal urgency and overcome executive inertia or pushback.",
            "Technical Integrity & Compliance (Architectural Readiness) — Demands technical depth over generic marketing speak. Needs confidence that the solution complies with regulatory standards (e.g., GDPR, HIPAA, SOC 2) and integrates with existing infrastructure without causing technical debt.",
        ],
    )

    add_persona(
        doc,
        "Technical Practitioner / Architect",
        "Cloud Architect, DevOps Engineer, SysAdmin, Site Reliability Engineer",
        "A hands-on builder and problem-solver responsible for implementing, optimizing, and maintaining core infrastructure and applications. They cut through high-level claims to focus on technical feasibility, operational stability, and efficiency. They attend events and engage with content to gain actionable knowledge, direct access to engineering experts, and tangible solutions that reduce toil, automate repetitive tasks, and keep mission-critical systems running predictably.",
        [
            "Proof Over Promise (Hands-on & Actionable) — Wants practical guidance to simplify complexity, standardize workflows, and automate routine tasks to reduce cognitive load and shift focus to high-value innovation.",
            "Operational Stability & Confidence (Consistency & Reliability) — Values a stable foundation across hybrid and multi-cloud environments. Relies on validated configurations, automated compliance checks and patching, and enterprise support to maintain reliability across mission-critical systems.",
            "Architectural Freedom (Openness & Interoperability) — Motivated by open-source technologies and platforms that run on the public cloud or hardware they choose, bridging siloes and providing a common foundation across dev and ops teams.",
            "Immediate Utility & Automation (Practical Impact) — Looks for built-in capabilities that cut manual steps in day-two operations, accelerate deployment pipelines, and proactively resolve issues before they cause outages.",
        ],
        alias="User briefs saying “Technical Practitioners & Influencers” usually map here (add Champion when advocacy/buying-urgency framing is required). Briefs saying Selectors (tactical choices about which technologies to acquire and use) also map here.",
    )

    add_persona(
        doc,
        "Developer",
        "Enterprise Software Engineer, Full-Stack Developer, Cloud-Native Developer, Application Architect",
        "A hands-on builder focused on writing clean code, solving complex technical problems, and shipping applications quickly. Developers value autonomy, low-overhead workflows, and modern open-source tools. They are deeply skeptical of marketing jargon and hyperbole. Instead, they seek authentic, code-first content, technical clarity, and active community engagement that helps them cut setup delays in their local development environment and CI/CD pipelines.",
        [
            "Code-First (Show, Don't Tell) — Wants to test and validate tools in their own environment immediately. Requires messaging that gets straight to the point, respecting their time by removing corporate fluff and leading with technical substance.",
            "Developer Experience & Productivity (Velocity) — Driven by speed, efficiency, and uninterrupted development. Responds to content that demonstrates how tools reduce boilerplate code, simplify containerization, and accelerate inner-loop development.",
            "Open-Source Authenticity & Peer-to-Peer Tone (Authentic Collaboration) — Values community-driven innovation and open standards. Demands a conversational, credible, and humble tone. They trust recommendations from fellow engineers and active community contributors rather than vendor marketing pitches.",
            "Practical Architecture & Modernization (Real-World Utility) — Needs practical guidance on integrating emerging technologies (like cloud-native frameworks, serverless, or local AI/LLM integration) into existing workflows without breaking their build or creating unmaintainable systems.",
        ],
    )

    heading(doc, "How to use this document", 2)
    bullets(
        doc,
        [
            "Upload to Google Drive, then Open with Google Docs if you want to edit online.",
            "Create or update one GenStudio Persona per heading. Paste Description and Messaging preferences into the persona fields. Roles can go in the description if the UI has no Roles field.",
            "This list is WIP and does not represent every persona configured in GenStudio.",
        ],
    )

    path = OUT / "GenStudio-Personas.docx"
    doc.save(path)
    print(f"Wrote {path.name}")


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    build_brand()
    build_products()
    build_personas()
